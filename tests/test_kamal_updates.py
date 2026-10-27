"""Regression tests for the Kamal grouping and optional order number updates."""

from pipelines.ordine_kamal import (
    OrdineKamalRow,
    assign_ordine_kamal_machines,
    build_ordine_kamal_rows,
)
from parsers.kamal_parser import KamalOrderRow


def _row(titolo: str, peso: float, coloredfm: str = "11") -> OrdineKamalRow:
    return OrdineKamalRow(coloredfm=coloredfm, titolo=titolo, peso_kg=peso)


def test_kamal_same_ply_titolo_values_are_combined_even_when_different():
    rows = [_row("50/1", 20), _row("96/1", 25)]

    assign_ordine_kamal_machines(rows)

    # Same colour and both /1: they are one 45 kg group, so both receive the
    # same machine assignment instead of two independent machines.
    assert [r.abbina for r in rows] == ["Machine 56 Cones", "Machine 56 Cones"]
    assert rows[0].macchina == rows[1].macchina


def test_kamal_odd_and_even_ply_titolo_values_never_share_machine():
    rows = [_row("50/1", 20), _row("100/2", 25)]

    assign_ordine_kamal_machines(rows)

    assert [r.abbina for r in rows] == ["Machine 24 Cones", "Machine 32 Cones"]
    assert rows[0].macchina != rows[1].macchina


def test_kamal_non_adjacent_different_titolo_stays_separate():
    rows = [_row("50/1", 20), _row("100/2", 20), _row("50/1", 20)]

    assign_ordine_kamal_machines(rows)

    # The two /1 rows share their own group, but /2 must not be merged into it.
    assert rows[0].abbina == rows[2].abbina == "Machine 56 Cones"
    assert rows[1].abbina == "Machine 24 Cones"


def test_kamal_optional_order_number_uses_requested_commento_format():
    parsed = [
        KamalOrderRow(
            reply_no="105",
            order_date="2026/07/20",
            coloredfm="11",
            titolo="50/1",
            peso_kg=10,
        )
    ]

    rows = build_ordine_kamal_rows(
        parsed,
        [{"articolo": "C1700001", "titolo": "50/1"}],
        order_no="28",
    )

    assert rows[0].commento == "PG-X-ORD.Nr.28-2026-(LETTER 105)"
    assert rows[0].order_no == "28"


def test_kamal_without_optional_order_number_keeps_legacy_commento():
    parsed = [KamalOrderRow(reply_no="105", order_date="2026/07/20")]
    rows = build_ordine_kamal_rows(parsed, [], order_no=None)
    assert rows[0].commento == "PG-X-ORD.nr.2026(LETTER 105)"


# ------------------------------------------------------- colour name -> DFM

from parsers.dfm_lookup import colour_name_key, lookup_dfm_color_by_name
from pipelines.ordine_kamal import resolve_kamal_colours

_DFM = [
    {"articolo": "C170115S", "coloredfm": "790351", "cldescr": "KA-MENT 4", "titolo": "80/2", "date": "2026-05-01"},
    {"articolo": "C170115S", "coloredfm": "790352", "cldescr": "KA-BEG 7 - 17-1210 TCX", "titolo": "80/2", "date": "2026-05-01"},
    {"articolo": "C170115S", "coloredfm": "790080", "cldescr": "KA- NIMBUA CLOUD 13-4108 TCX", "titolo": "80/2", "date": "2026-05-01"},
    {"articolo": "C170115S", "coloredfm": "790266", "cldescr": "KA-FUNGI 99(17-1212 TCX)", "titolo": "80/2", "date": "2026-05-01"},
]


def test_colour_name_key_ignores_ka_tag_pantone_arabic_and_spacing():
    spellings = ["KA-MENT 4", "KA-MENT 4 - 13-0608 TCX", "MENT 4", "صMENT 4", "ment  4", "MENT 4-KA", "KA-MENT4"]
    assert {colour_name_key(s) for s in spellings} == {"MENT4"}
    assert colour_name_key("KA-FUNGI 99(17-1212 TCX)") == colour_name_key("FUNGI 99") == "FUNGI99"
    assert colour_name_key("KA- NIMBUA CLOUD 13-4108 TCX") == "NIMBUACLOUD"
    assert colour_name_key("") == colour_name_key("   ") == ""
    # numbers that are part of the name stay significant
    assert colour_name_key("KA-GRY 10") != colour_name_key("KA-GRY 1")


def test_lookup_by_name_finds_the_dfm_entry_whichever_spelling_either_side_uses():
    assert lookup_dfm_color_by_name("صMENT 4", _DFM) == ("790351", "KA-MENT 4")
    assert lookup_dfm_color_by_name("MENT 4 - 13-0608 TCX", _DFM) == ("790351", "KA-MENT 4")
    assert lookup_dfm_color_by_name("BEG 7", _DFM) == ("790352", "KA-BEG 7 - 17-1210 TCX")
    assert lookup_dfm_color_by_name("NIMBUA CLOUD", _DFM) == ("790080", "KA- NIMBUA CLOUD 13-4108 TCX")
    assert lookup_dfm_color_by_name("FUNGI 99", _DFM) == ("790266", "KA-FUNGI 99(17-1212 TCX)")
    assert lookup_dfm_color_by_name("NEVER DYED", _DFM) == ("", "")
    assert lookup_dfm_color_by_name("", _DFM) == ("", "")


def test_lookup_by_name_prefers_the_same_articolo_then_the_newest_entry():
    entries = [
        {"articolo": "C170001S", "coloredfm": "1", "cldescr": "KA-MENT 4", "date": "2026-09-01"},
        {"articolo": "C170115S", "coloredfm": "2", "cldescr": "KA-MENT 4", "date": "2026-01-01"},
        {"articolo": "C170115S", "coloredfm": "3", "cldescr": "KA-MENT 4 - 13-0608 TCX", "date": "2026-03-01"},
    ]
    assert lookup_dfm_color_by_name("MENT 4", entries, articolo="C170115S")[0] == "3"   # same article, newest
    assert lookup_dfm_color_by_name("MENT 4", entries)[0] == "1"                        # no article: newest overall


def test_ordine_kamal_takes_code_and_name_from_dfm_or_keeps_the_pdf_name_with_blank_code():
    parsed = [
        KamalOrderRow(titolo="80/2", colore_name="صMENT 4", peso_kg=10),
        KamalOrderRow(titolo="80/2", colore_name="BEG 7", coloredfm="", peso_kg=10),
        KamalOrderRow(titolo="80/2", colore_name="BRAND NEW", peso_kg=10),
    ]
    rows = build_ordine_kamal_rows(parsed, _DFM)
    assert [r.coloredfm for r in rows] == ["790351", "790352", ""]
    assert [r.colore_name for r in rows] == ["KA-MENT 4", "KA-BEG 7 - 17-1210 TCX", "BRAND NEW"]
    # the parsed rows (the raw sheet's source) were updated the same way, and a second pass changes nothing
    assert [r.coloredfm for r in parsed] == ["790351", "790352", ""]
    assert resolve_kamal_colours(parsed, _DFM) == 2
    assert [r.coloredfm for r in parsed] == ["790351", "790352", ""]


def test_kamal_workbook_raw_sheet_shows_the_dfm_code_and_name(tmp_path):
    import openpyxl
    from exporters.kamal_excel_exporter import KamalExcelExporter

    out = tmp_path / "k.xlsx"
    rows = [KamalOrderRow(titolo="80/2", colore_name="صMENT 4", peso_kg=10, order_date="2026/10/20", reply_no="108"),
            KamalOrderRow(titolo="80/2", colore_name="BRAND NEW", peso_kg=10, order_date="2026/10/20", reply_no="108")]
    KamalExcelExporter(rows, out, dfm_c170_entries=_DFM).export()
    wb = openpyxl.load_workbook(out)
    raw = [[c.value for c in r] for r in wb["Kamal Order"].iter_rows(min_row=1)]
    header = raw[0]
    cf, cn = header.index("COLOREDFM"), header.index("Colore")
    assert (raw[1][cf], raw[1][cn]) == ("790351", "KA-MENT 4")
    assert (raw[2][cf] or "", raw[2][cn]) == ("", "BRAND NEW")
    ordine = [[c.value for c in r] for r in wb["Ordine Kamal"].iter_rows(min_row=1)]
    assert ordine[1][ordine[0].index("COLORE")] in (790351, "790351")


# --------------------------------------------- Ordine Kamal workbook -> Create (EXCEL+Biglietti)

def test_the_kamal_tab_workbook_loads_from_its_order_sheet_not_the_raw_pdf_sheet(tmp_path):
    """The Kamal-tab workbook lists the raw PDF sheet ("Kamal Order") first; the order
    lines must come from "Ordine Kamal", with the colour NAME taken from the PDF sheet."""
    from exporters.biglietti_exporter import detect_order_format, load_el_kamal_order
    from exporters.kamal_excel_exporter import KamalExcelExporter

    out = tmp_path / "kamal.xlsx"
    rows = [KamalOrderRow(titolo="80/2", colore_name="صMENT 4", peso_kg=10, order_date="2026/10/20", reply_no="108",
                          notes="PG-1-ORD.Nr.28-2026(LETTER 108)"),
            KamalOrderRow(titolo="80/2", colore_name="BEG 7", peso_kg=20.5, order_date="2026/10/20", reply_no="108",
                          notes="PG-2-ORD.Nr.28-2026(LETTER 108)")]
    KamalExcelExporter(rows, out, dfm_c170_entries=_DFM).export()

    assert detect_order_format(out) == "EL_KAMAL"
    records, raw_rows = load_el_kamal_order(out)
    assert raw_rows == []
    assert [r.article for r in records] == ["C170115S", "C170115S"]
    assert [r.color_code for r in records] == ["790351", "790352"]
    assert [r.color_name for r in records] == ["KA-MENT 4", "KA-BEG 7 - 17-1210 TCX"]
    assert [r.kg for r in records] == [10, 20.5]                 # Q.TA is kilograms: not divided by 10
    assert all(r.quantity_cones is None for r in records)        # and it is not a cone count
    assert all("ORD" in r.commento for r in records)
