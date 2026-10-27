"""Ordine Elvy's Prezzi-PO file and the ELVY file names of Create (EXCEL+Biglietti)."""
from dataclasses import replace

import openpyxl
import pytest

from exporters.biglietti_exporter import build_output_stem
from parsers.pdf_parser import OrderRow
from pipelines.prezzi_elvy import export_prezzi_po, machine_capacity, po_label, prezzi_rows


def _line(**kw):
    base = dict(
        pod_no="1500", pos=10, articolo_delta="G010937S", coloredfm="343", cldescr="AZZURRINO",
        colour="AZZURRINO", ne="120/2", quantity_cones=72, abbina="Machine 72 Cones",
        livello=60, prezzo=4.59,
    )
    base.update(kw)
    return OrderRow(**base)


# ------------------------------------------------------------------ Prezzi-PO
def test_machine_capacity_comes_from_the_abbina():
    assert machine_capacity("Machine 72 Cones") == 72
    assert machine_capacity("machine 24 cone") == 24
    assert machine_capacity("") is None and machine_capacity("no machine") is None


def test_po_label_lists_each_po_once_from_the_pod_no_column():
    assert po_label([_line(pod_no="1500"), _line(pod_no="1500")]) == "1500"
    assert po_label([_line(pod_no="1500"), _line(pod_no="PO-1501")]) == "1500 & 1501"
    assert po_label([_line(pod_no="")]) == ""


def test_prezzi_rows_add_two_dollars_only_for_the_surcharge_machines():
    rows = prezzi_rows(
        [_line(abbina="Machine 72 Cones", prezzo=4.59), _line(pos=20, abbina="Machine 24 Cones", prezzo=6.16, coloredfm="895"),
         _line(pos=30, abbina="Machine 32 Cones", prezzo=None, coloredfm="306")],
        {"C010937S": "120/2 ORGANIC"},
    )
    # Riga, ARTICOLO, TITOLO, COLORE, Q.TA, DESCR COL, LIVELLO, M/C, PREZZO, prezzo + (2$)
    assert rows[0] == [10, "C010937S", "120/2 ORGANIC", "343", 72, "AZZURRINO", 60, 72, 4.59, 4.59]
    assert rows[1][7:] == [24, 6.16, 8.16]                      # 24 Rocche machine: +2$
    assert rows[2][7:] == [32, None, None]                      # no price found: empty cells


def test_titolo_falls_back_to_the_pdfs_ne():
    assert prezzi_rows([_line()], {})[0][2] == "120/2"


def test_export_prezzi_po_names_the_file_after_the_po(tmp_path):
    path = export_prezzi_po(tmp_path, [_line(), _line(pod_no="1501", pos=20, coloredfm="372")], {})
    assert path.name == "Prezzi-PO-1500 & 1501.xlsx"
    ws = openpyxl.load_workbook(path)["ELVY"]
    assert [c.value for c in ws[1]] == [
        "Riga", "ARTICOLO", "TITOLO", "COLORE", "Q.TA", "DESCR COL", "LIVELLO", "M/C", "PREZZO", "prezzo + (2$)",
    ]
    assert ws.max_row == 3
    assert export_prezzi_po(tmp_path, [], {}) is None
    assert export_prezzi_po(tmp_path, [_line(pod_no="")], {}).name == "Prezzi-PO.xlsx"


# ------------------------------------------------------------------ ELVY file names
def _elvy(comment):
    from tests.test_med_commento_prezzi import _shared_record
    return replace(_shared_record("3009", ""), additional_raw=comment)


@pytest.mark.parametrize("comments, expected", [
    (["PG-X-PO-1500-2026"], "PO-1500-2026"),
    (["PG-X-PO-1500-2026", "PG-157474-PO-1500-2026"], "PO-1500-2026"),              # the same one repeated
    (["PG-X-PO-1500-2026", "PG-X-PO-1501-2026", "PG-X-PO-1500-2026"], "PO-1500 & 1501-2026"),
    (["PG-X-PO-1500-2026", "PG-X-PO-90-2025"], "PO-1500-2026 & PO-90-2025"),         # different years stay apart
    (["PG-X-PO-1500-2026-ref.ABC"], "PO-1500-2026"),
    (["PG-X-PO1500-2026"], "PO-1500-2026"),
    (["PG-X", ""], "ELVY"),
])
def test_elvy_files_are_named_from_the_commento_values(comments, expected):
    assert build_output_stem([_elvy(c) for c in comments], "ELVY") == expected


def test_word_files_no_longer_carry_the_word_biglietti():
    source = open("gui/tabs/biglietti_tab.py", encoding="utf-8").read()
    assert "_Biglietti.docx" not in source and "_ELVY_Biglietti" not in source and "_MED_Biglietti" not in source
