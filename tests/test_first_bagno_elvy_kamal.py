"""First Bagno for Ordine Elvy / Ordine Kamal, typed ERP columns, related fixes."""
from datetime import datetime
from unittest.mock import patch

import openpyxl
import pytest

from calculate.bagno import is_valid_first_bagno, parse_first_bagno, sequence_bagno


# ---------------------------------------------------------------- helper
def test_sequence_bagno_first_row_blank_following_rows_carry_the_bath():
    keys = ["a", "a", "b", "c", "c", "c"]
    assert sequence_bagno(keys, "T1200") == ["", "T1200", "", "", "T1202", "T1202"]


def test_sequence_bagno_non_consecutive_rows_of_one_group_share_the_bath():
    # Kamal groups by colour wherever the rows sit in the letter.
    assert sequence_bagno(["a", "b", "a", "b"], "T10") == ["", "", "T10", "T11"]


def test_sequence_bagno_rows_without_a_group_each_take_one_number():
    assert sequence_bagno([None, "a", None, "a"], "S5") == ["", "", "", "S6"]


def test_sequence_bagno_disabled_when_empty_and_rejects_bad_format():
    assert sequence_bagno(["a", "a"], None) == ["", ""]
    assert sequence_bagno(["a", "a"], "  ") == ["", ""]
    with pytest.raises(ValueError):
        sequence_bagno([], "1200")
    with pytest.raises(ValueError):
        sequence_bagno(["a"], "T-12")
    assert parse_first_bagno(" s1200 ") == ("s", 1200)
    assert is_valid_first_bagno("T1200") and not is_valid_first_bagno("1200")


# ---------------------------------------------------------------- Elvy
def _order_row(colour, qty, abbina="", pod="POD1", articolo="G130027S", dfm="440"):
    from parsers.pdf_parser import OrderRow
    return OrderRow(
        pod_no=pod, po_date="19/4/2026", articolo_delta=articolo, coloredfm=dfm,
        colour=colour, quantity_cones=qty, abbina=abbina,
    )


def test_elvy_bagno_follows_same_colour_runs_and_machine_groups():
    from calculate.abbina_calculator import AbbinaCalculator
    from pipelines.ordini_elvy import build_ordini_elvy_rows

    rows = [
        _order_row("100", 24), _order_row("200", 20), _order_row("200", 10),
        _order_row("300", 6),
    ]
    AbbinaCalculator().calculate(rows)
    out = build_ordini_elvy_rows(rows, first_bagno="T1200")
    assert [r.bagno_proposto for r in out] == ["", "", "T1201", ""]
    # the two 200 rows are one machine batch (30 cones -> 32) ...
    assert out[1].macchina == out[2].macchina
    # ... and that is exactly the group that shares the bath
    assert out[1].bagno_group == out[2].bagno_group != out[0].bagno_group


def test_elvy_other_po_never_joins_a_run_and_default_is_blank():
    from calculate.abbina_calculator import AbbinaCalculator
    from pipelines.ordini_elvy import build_ordini_elvy_rows

    rows = [_order_row("100", 24, pod="A"), _order_row("100", 24, pod="B")]
    for r in rows:
        AbbinaCalculator().calculate([r])
    assert [r.bagno_proposto for r in build_ordini_elvy_rows(rows, first_bagno="T1")] == ["", ""]
    assert [r.bagno_proposto for r in build_ordini_elvy_rows(rows)] == ["", ""]


def test_elvy_dates_are_midnight_real_dates():
    from pipelines.ordini_elvy import build_ordini_elvy_rows

    row = build_ordini_elvy_rows([_order_row("100", 24)])[0]
    for value in (row.data_consegna, row.data_riconsegna):
        assert isinstance(value, datetime)
        assert (value.hour, value.minute, value.second, value.microsecond) == (0, 0, 0, 0)


def test_elvy_erp_file_has_typed_columns_and_bagno_text(tmp_path):
    from calculate.abbina_calculator import AbbinaCalculator
    from pipelines.ordini_elvy import build_ordini_elvy_rows, export_ordini_full

    rows = [_order_row("200", 20), _order_row("200", 10)]
    AbbinaCalculator().calculate(rows)
    target = tmp_path / "erp.xlsx"
    export_ordini_full(target, build_ordini_elvy_rows(rows, first_bagno="T1200"))
    ws = openpyxl.load_workbook(target)["ORDINE VENDITA EGITTO"]
    header = [c.value for c in ws[1]]
    second = {h: ws.cell(row=3, column=i + 1) for i, h in enumerate(header)}
    assert second["BAGNO PREPOSTO"].value == "T1200"
    for name in ("CLIENTE", "COLORE", "Q.TA", "LAVORANTE", "LAV. SUCC", "MAG. GREGGIO", "GRUPPO MACCHINA"):
        assert isinstance(second[name].value, int), name
    for name in ("CONSEGNA", "DATA RICONSEGNA"):
        assert isinstance(second[name].value, datetime), name


def test_elvy_workbook_sheet_gets_the_same_bagno(tmp_path):
    from exporters.excel_exporter import ExcelExporter

    rows = [_order_row("200", 20), _order_row("200", 10)]
    ExcelExporter(rows, tmp_path / "o.xlsx", first_bagno="T1200").export()
    ws = openpyxl.load_workbook(tmp_path / "o.xlsx")["Ordini ELVY"]
    header = [c.value for c in ws[1]]
    col = header.index("BAGNO PREPOSTO") + 1
    assert [ws.cell(row=r, column=col).value for r in (2, 3)] in (["", "T1200"], [None, "T1200"])


def test_elvy_invalid_first_bagno_is_rejected():
    from pipelines.ordini_elvy import build_ordini_elvy_rows

    with pytest.raises(ValueError):
        build_ordini_elvy_rows([_order_row("100", 24)], first_bagno="1200")


# ---------------------------------------------------------------- Kamal
def _kamal(coloredfm, kg, colore="", reply="98"):
    from pipelines.ordine_kamal import OrdineKamalRow
    return OrdineKamalRow(coloredfm=coloredfm, peso_kg=kg, colore_name=colore, articolo_delta="C1700001")


def test_kamal_bagno_one_bath_per_colour_even_when_rows_are_apart():
    from pipelines.ordine_kamal import assign_bagno_prepostato, assign_ordine_kamal_machines

    rows = [_kamal("11", 10), _kamal("22", 8), _kamal("11", 12), _kamal("", 5, colore="Navy"),
            _kamal("", 5, colore="Navy"), _kamal("", 3)]
    assign_ordine_kamal_machines(rows)
    assign_bagno_prepostato(rows, "T1200")
    # 11 -> T1200 (rows 0,2) | 22 -> T1201 | Navy -> T1202 (rows 3,4) | lone row -> T1203
    assert [r.bagno_proposto for r in rows] == ["", "", "T1200", "", "T1202", ""]
    # same colour -> same machine, as before
    assert rows[0].macchina == rows[2].macchina


def test_kamal_bagno_blank_when_disabled_and_validated_when_given():
    from pipelines.ordine_kamal import assign_bagno_prepostato

    rows = [_kamal("11", 10), _kamal("11", 10)]
    assign_bagno_prepostato(rows, "")
    assert [r.bagno_proposto for r in rows] == ["", ""]
    with pytest.raises(ValueError):
        assign_bagno_prepostato(rows, "abc")


def test_kamal_workbook_end_to_end_with_text_colour_code_and_bagno(tmp_path):
    from exporters.kamal_excel_exporter import KamalExcelExporter
    from parsers.kamal_parser import KamalOrderRow

    parsed = [
        KamalOrderRow(reply_no="98", order_date="2026/07/20", coloredfm="11", peso_kg=10.5),
        KamalOrderRow(reply_no="98", order_date="2026/07/20", coloredfm="", colore_name="Navy", peso_kg=4),
        KamalOrderRow(reply_no="98", order_date="2026/07/20", coloredfm="11", peso_kg=2),
        KamalOrderRow(reply_no="98", order_date="2026/07/20", coloredfm="A12", peso_kg=2),  # non-numeric
    ]
    out = tmp_path / "k.xlsx"
    KamalExcelExporter(parsed, out, first_bagno="T50").export()
    wb = openpyxl.load_workbook(out)
    ws = wb["Ordine Kamal"]
    header = [c.value for c in ws[1]]
    bagno = [ws.cell(row=r, column=header.index("BAGNO PREPOSTO") + 1).value for r in range(2, 6)]
    assert [b or "" for b in bagno] == ["", "", "T50", ""]
    colore = [ws.cell(row=r, column=header.index("COLORE") + 1).value for r in range(2, 6)]
    assert colore[0] == 11 and colore[3] == "A12"          # text code no longer crashes the export
    # raw sheet: Abbina is now also filled for the row grouped by Colore (blank COLOREDFM)
    raw = wb["Kamal Order"]
    raw_header = [c.value for c in raw[1]]
    abbina = [raw.cell(row=r, column=raw_header.index("Abbina") + 1).value for r in range(2, 6)]
    assert all(a for a in abbina)


# ---------------------------------------------------------------- Filato reader
def test_filato_reader_missing_sheet_is_ok_for_extraction(tmp_path):
    from pipelines.ordini_elvy import read_filato_tinturia_sheet

    path = tmp_path / "no_filato.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Purchase Orders"
    wb.save(path)
    assert read_filato_tinturia_sheet(path, missing_ok=True) == []
    with pytest.raises(ValueError):
        read_filato_tinturia_sheet(path)


# ---------------------------------------------------------------- Med system / ERP sheet types
def _med_record(**kw):
    from pipelines.ordine_med import OrdineMedRow
    base = dict(
        riga=1, code_org="C130027S", titolo="30/1", descr_col="Blue", articolo="C130027S",
        colore="304", rocc=24, abbin=None, consegna_input=None, pt_grg="", pt_med="",
        polmoni="", cliente_note="", nota_grg="", nota_col="", kg_note="", fabb="", prezz_note="",
        cliente="3009", gruppo_macchina=3302, consegna=datetime(2026, 10, 20),
        data_riconsegna=datetime(2026, 10, 19), commento="PG-X", bagno_prepostato="T1201",
    )
    base.update(kw)
    return OrdineMedRow(**base)


def test_med_erp_sheet_columns_have_real_types(tmp_path):
    from pipelines.ordine_med import export_erp_order_workbook

    target = tmp_path / "erp.xlsx"
    export_erp_order_workbook(target, [_med_record(), _med_record(colore="A12")])
    ws = openpyxl.load_workbook(target)["Dati sistema (B-N)"]
    header = [c.value for c in ws[1]]
    assert "SIGLA DISPOSIZIONE" in header and "SIGLA DISPOSIZONE" not in header
    cell = {h: ws.cell(row=2, column=i + 1) for i, h in enumerate(header)}
    for name in ("CLIENTE", "COLORE", "Q.TA", "LAVORANTE", "LAV. SUCC", "MAG. GREGGIO", "GRUPPO MACCHINA"):
        assert isinstance(cell[name].value, int), name
    assert cell["CLIENTE"].value == 3009 and cell["LAVORANTE"].value == 900901
    for name in ("CONSEGNA", "DATA RICONSEGNA"):
        assert isinstance(cell[name].value, datetime), name
        assert cell[name].number_format == "dd/mm/yyyy"
    assert cell["BAGNO PREPOSTO"].value == "T1201" and cell["BAGNO PREPOSTO"].number_format == "@"
    assert ws.cell(row=3, column=header.index("COLORE") + 1).value == "A12"   # non-numeric stays text


def test_med_consegna_text_dates_become_real_dates():
    from pipelines.ordine_med import _to_date_value, compute_data_riconsegna

    assert _to_date_value("20/10/2026") == datetime(2026, 10, 20)
    assert _to_date_value(datetime(2026, 10, 20, 15, 30)) == datetime(2026, 10, 20)
    assert _to_date_value("") is None and _to_date_value("soon") == "soon"
    record = _med_record(consegna=None, consegna_input=_to_date_value("20/10/2026"))
    compute_data_riconsegna([record])
    assert record.data_riconsegna == datetime(2026, 10, 19)


# ---------------------------------------------------------------- GUI wiring
tk = pytest.importorskip("tkinter")


@pytest.fixture
def root():
    try:
        window = tk.Tk()
    except tk.TclError:
        pytest.skip("no display available")
    window.withdraw()
    yield window
    window.destroy()


def test_kamal_tab_has_first_bagno_controls_and_rejects_bad_value(root, monkeypatch, tmp_path):
    from gui.tabs import kamal_tab

    monkeypatch.setattr(kamal_tab, "load_settings", lambda: {})
    monkeypatch.setattr(kamal_tab, "save_settings", lambda prefs: None)
    tab = kamal_tab.KamalTab(root)
    tab._pdf_paths = [tmp_path / "x.pdf"]
    tab._output_path = tmp_path / "o.xlsx"
    tab._shared_dfm_path = "dfm"
    tab._kamal_bagno_enabled.set(True)
    tab._kamal_bagno_start.set("1200")
    with patch.object(kamal_tab.messagebox, "showwarning") as warn, \
            patch.object(kamal_tab.threading, "Thread") as thread:
        tab._on_convert()
    assert warn.called and not thread.called           # bad format stops before any work
    tab._kamal_bagno_start.set("T1200")
    with patch.object(kamal_tab.messagebox, "showwarning") as warn, \
            patch.object(kamal_tab.threading, "Thread") as thread:
        tab._on_convert()
    assert not warn.called and thread.called and tab._first_bagno_snapshot == "T1200"


def test_main_window_source_wires_first_bagno_for_elvy():
    from pathlib import Path
    root_dir = Path(__file__).resolve().parent.parent
    po = (root_dir / "gui" / "workflows" / "purchase_orders.py").read_text(encoding="utf-8")
    app = (root_dir / "gui" / "gui.py").read_text(encoding="utf-8")
    assert "Bagno Preposato (ELVY)" in po and po.count("first_bagno=first_bagno") >= 3
    assert "_po_bagno_enabled" in app and "_po_bagno_start" in app


def test_kamal_tab_conversion_uses_sqlite_stock_and_lotti_and_first_bagno(root, monkeypatch, tmp_path):
    """The SQLite snapshots (not only a freshly read Excel) must feed raw-yarn
    matching: Lotto first, stock second -- and First bagno reaches both files."""
    import pandas as pd
    from gui.tabs import kamal_tab
    from parsers.kamal_parser import KamalOrderRow

    monkeypatch.setattr(kamal_tab, "load_settings", lambda: {})
    monkeypatch.setattr(kamal_tab, "save_settings", lambda prefs: None)
    tab = kamal_tab.KamalTab(root)

    rows = [
        KamalOrderRow(reply_no="98", order_date="2026/07/20", sender_msg_no="MSG1", coloredfm="11", peso_kg=10),
        KamalOrderRow(reply_no="98", order_date="2026/07/20", sender_msg_no="MSG1", coloredfm="11", peso_kg=5),
    ]

    class FakeParser:
        def __init__(self, path):
            pass

        def parse(self):
            return rows

    entries = [{"articolo": "C1700001", "titolo": "80/2"}]
    monkeypatch.setattr(kamal_tab, "KamalParser", FakeParser)
    monkeypatch.setattr(kamal_tab, "load_dfm_entries_by_prefix", lambda prefix: entries)
    stock = pd.DataFrame([{"articolo": "G1700001", "titolo": "80/2", "partita": "777", "mag_rocche": 50, "mag_peso": 500.0}])
    lotti = pd.DataFrame([{"articolo": "G1700001", "partita": "888", "ordine": 0, "qesi": 5, "lotto": "MSG1", "magazzino": 900910}])
    monkeypatch.setattr(kamal_tab.magazino_logic, "load_magazino_snapshot", lambda prefix=None: (stock, []))
    monkeypatch.setattr(kamal_tab.lotti_logic, "load_lotti_snapshot", lambda prefix=None: (lotti, []))

    erp_dir, filato_dir = tmp_path / "erp", tmp_path / "filato"
    erp_dir.mkdir()
    filato_dir.mkdir()
    tab._pdf_paths = [tmp_path / "x.pdf"]
    tab._output_path = tmp_path / "Ordine_Kamal.xlsx"
    tab._kamal_erp_export_dir, tab._kamal_filato_export_dir = erp_dir, filato_dir
    tab._first_bagno_snapshot = "T1200"
    tab.after = lambda *a, **k: None          # no Tk main loop in this test
    tab._set_status = lambda text: None
    tab._run_conversion()

    erp = openpyxl.load_workbook(erp_dir / "EXCEL PER ORDINE VENDITA EGITTO.xlsx")["ORDINE VENDITA EGITTO"]
    header = [c.value for c in erp[1]]
    comment = [erp.cell(row=r, column=header.index("COMMENTO") + 1).value for r in (2, 3)]
    bagno = [erp.cell(row=r, column=header.index("BAGNO PREPOSTO") + 1).value for r in (2, 3)]
    assert all("PG-888" in c for c in comment)          # matched by Lotto from the SQLite snapshot
    assert [b or "" for b in bagno] == ["", "T1200"]
    assert (filato_dir / "Filato x Tinturia.xlsx").is_file()


# ------------------------------------------------ oversized colour groups are split
def test_split_into_batches_never_overfills_the_largest_machine():
    from calculate.abbina_calculator import split_into_batches
    assert split_into_batches([400, 400]) == [0, 1]
    assert split_into_batches([300, 300, 100, 50]) == [0, 0, 1, 1]   # 600 + 100 would pass 672
    assert split_into_batches([300, 300, 72]) == [0, 0, 0]       # exactly 672 still fits
    assert split_into_batches([672]) == [0]
    assert split_into_batches([]) == []


def test_elvy_same_colour_over_672_goes_to_two_machines_and_two_baths():
    from calculate.abbina_calculator import AbbinaCalculator
    from pipelines.ordini_elvy import build_ordini_elvy_rows

    rows = [_order_row("100", 400), _order_row("100", 400), _order_row("100", 20), _order_row("200", 24)]
    AbbinaCalculator().calculate(rows)
    assert [r.abbina for r in rows] == [
        "Machine 672 Cones", "Machine 672 Cones", "Machine 672 Cones", "Machine 24 Cones"]
    out = build_ordini_elvy_rows(rows, first_bagno="T1")
    # batch 0 = row 0 (T1), batch 1 = rows 1+2 (T2), colour 200 = T3
    assert [r.bagno_proposto for r in out] == ["", "", "T2", ""]


def test_elvy_different_colour_codes_never_share_a_machine():
    from calculate.abbina_calculator import AbbinaCalculator
    rows = [_order_row("007", 10), _order_row("7", 10)]
    AbbinaCalculator().calculate(rows)
    assert rows[0].abbina_batch == 0 and rows[1].abbina_batch == 0
    assert [r.abbina for r in rows] == ["Machine 24 Cones"] * 2


def test_kamal_same_colour_over_largest_machine_is_split():
    from pipelines.ordine_kamal import assign_bagno_prepostato, assign_ordine_kamal_machines

    rows = [_kamal("11", 400), _kamal("11", 300), _kamal("11", 30)]
    assign_ordine_kamal_machines(rows)
    # 400 alone -> 672 | 300 + 30 = 330 -> the smallest machine that holds it (384)
    assert [r.abbina for r in rows] == ["Machine 672 Cones", "Machine 384 Cones", "Machine 384 Cones"]
    assert [r.batch_no for r in rows] == [0, 1, 1]
    assign_bagno_prepostato(rows, "T1")
    assert [r.bagno_proposto for r in rows] == ["", "", "T2"]


def test_partial_pdf_abbina_fills_blank_rows_of_the_same_colour_run():
    from calculate.abbina_calculator import AbbinaCalculator
    from pipelines.ordini_elvy import build_ordini_elvy_rows

    rows = [_order_row("100", 10), _order_row("100", 10, abbina="Machine 32 Cones"), _order_row("100", 10)]
    AbbinaCalculator().calculate(rows, only_if_missing=True)
    assert [r.abbina for r in rows] == ["Machine 32 Cones"] * 3
    out = build_ordini_elvy_rows(rows, first_bagno="T1")
    assert len({r.macchina for r in out}) == 1 and [r.bagno_proposto for r in out] == ["", "T1", "T1"]


# ------------------------------------------------ rename Planing -> Planning
def test_app_data_folder_migrates_from_the_old_misspelt_name(tmp_path, monkeypatch):
    import sys
    from utility import utils

    (tmp_path / "Planing" / "settings").mkdir(parents=True)
    (tmp_path / "Planing" / "settings" / "prefs.json").write_text("{}", encoding="utf-8")
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))

    new_dir = utils._get_app_data_dir()
    assert new_dir == tmp_path / "Planning"
    assert (new_dir / "settings" / "prefs.json").is_file() and not (tmp_path / "Planing").exists()


def test_app_data_migration_merges_when_new_folder_already_exists(tmp_path, monkeypatch):
    import sys
    from utility import utils

    (tmp_path / "Planing" / "logs").mkdir(parents=True)
    (tmp_path / "Planing" / "logs" / "converter.log").write_text("old", encoding="utf-8")
    (tmp_path / "Planing" / "startup_error.log").write_text("old error", encoding="utf-8")
    (tmp_path / "Planning").mkdir()
    (tmp_path / "Planning" / "startup_error.log").write_text("new error", encoding="utf-8")   # already there
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))

    new_dir = utils._get_app_data_dir()
    assert (new_dir / "logs" / "converter.log").read_text(encoding="utf-8") == "old"
    assert (new_dir / "startup_error.log").read_text(encoding="utf-8") == "new error"   # never overwritten


def test_no_misspelt_name_left_in_source_except_deliberate_compat():
    from pathlib import Path
    root = Path(__file__).resolve().parent.parent
    allowed = {"utility/utils.py", "utility/updater.py", "installer.iss", "build.bat",
               "tests/test_first_bagno_elvy_kamal.py"}
    offenders = []
    for path in root.rglob("*"):
        if path.suffix not in {".py", ".bat", ".spec", ".iss", ".md"} or ".git" in path.parts:
            continue
        rel = path.relative_to(root).as_posix()
        if rel not in allowed and "Planing" in path.read_text(encoding="utf-8", errors="ignore"):
            offenders.append(rel)
    assert offenders == []
