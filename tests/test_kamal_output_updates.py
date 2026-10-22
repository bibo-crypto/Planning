from pathlib import Path

import openpyxl

from exporters.biglietti_exporter import (
    OrderRecord,
    _kamal_kg,
    build_el_kamal_excel_stem,
    build_el_kamal_stem,
    export_workbook,
)


def _record(commento="PG-159640-ORD.Nr.28-2026(LETTER 108)"):
    return OrderRecord(
        customer_code="3019",
        customer_name="EL KAMAL",
        article="C170026S",
        description="Blue",
        additional_raw=commento,
        color_code="790001",
        color_name="Blue",
        order_no="28",
        order_row="1",
        colored_batch="159640",
        raw_batch="159640",
        quantity_cones=56,
        raw_weight=52.5,
        dispo="D-00505582-001",
        bagno="T1303",
        delivery="20/10/2026",
        machine="3",
        title="50/1",
        commento=commento,
        kg=52.5,
    )


def test_kamal_output_stems_follow_requested_names():
    records = [_record()]
    assert build_el_kamal_excel_stem(records) == "ORD.Nr.28-2026(LETTER 108)"
    assert build_el_kamal_stem(records) == "D-505582-KAMAL-2026"


def test_kamal_kg_converts_erp_tenths_without_touching_decimal_values():
    assert _kamal_kg(525) == 52.5
    assert _kamal_kg("525") == 52.5
    assert _kamal_kg("52.5") == 52.5
    assert _kamal_kg("52,5") == 52.5


def test_kamal_workbook_can_hide_rocche_and_kg(tmp_path):
    path = tmp_path / "kamal.xlsx"
    record = _record()
    record.vmm22 = 49.25
    export_workbook(path, [record], [], customer="EL_KAMAL", hide_kamal_quantities=True)
    ws = openpyxl.load_workbook(path, data_only=True)["EL_KAMAL"]
    header = [cell.value for cell in ws[1]]
    row = {header[i]: ws.cell(row=2, column=i + 1).value for i in range(len(header))}
    assert row["Rocche"] is None and row["KG"] is None
    assert "VMM22" not in row                       # EL KAMAL has no VMM22 column

    export_workbook(path, [record], [], customer="EL_KAMAL", hide_kamal_quantities=False)
    ws = openpyxl.load_workbook(path, data_only=True)["EL_KAMAL"]
    header = [cell.value for cell in ws[1]]
    row = {header[i]: ws.cell(row=2, column=i + 1).value for i in range(len(header))}
    assert row["Rocche"] == 56 and row["KG"] == 52.5


def test_erp_and_filato_exports_do_not_leave_previous_workbooks(tmp_path):
    from pipelines.ordini_elvy import OrdiniElvyRow, RawYarnMatch, export_filato_full, export_ordini_full

    erp_path = tmp_path / "EXCEL PER ORDINE VENDITA EGITTO.xlsx"
    filato_path = tmp_path / "Filato x Tinturia.xlsx"
    export_ordini_full(erp_path, [OrdiniElvyRow(articolo_delta="C130001S", coloredfm="1")])
    export_ordini_full(erp_path, [OrdiniElvyRow(articolo_delta="C130002S", coloredfm="2")])
    export_filato_full(filato_path, [RawYarnMatch(articolo="G130001S", partita="1")])
    export_filato_full(filato_path, [RawYarnMatch(articolo="G130002S", partita="2")])

    assert not (tmp_path / "EXCEL PER ORDINE VENDITA EGITTO.previous.xlsx").exists()
    assert not (tmp_path / "Filato x Tinturia.previous.xlsx").exists()


def test_the_create_page_never_deletes_earlier_orders():
    """A new order must not wipe the files of earlier orders -- for any customer."""
    source = Path("gui/tabs/biglietti_tab.py").read_text(encoding="utf-8")
    assert "_clear_previous" not in source and ".unlink(" not in source and "rmtree" not in source
