"""Fixes that came out of the external review report (each one reproduced first)."""
import os
import sqlite3
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
import pytest


def _workbook(path: Path, sheets: dict[str, list[list]]) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    wb.save(path)


# ------------------------------------------------------------ 1. atomic Excel save
def test_failed_save_leaves_the_original_workbook_untouched(tmp_path):
    from utility.excel_io import safe_save_workbook

    target = tmp_path / "orders.xlsx"
    _workbook(target, {"Orders": [["Partita Col"], ["A1"]]})
    original = target.read_bytes()

    class DiskFull:
        def save(self, destination):
            Path(destination).write_bytes(b"PK\x03\x04 half written")
            raise OSError(28, "No space left on device")

    with pytest.raises(OSError):
        safe_save_workbook(DiskFull(), target)
    assert target.read_bytes() == original
    assert [p.name for p in tmp_path.iterdir()] == ["orders.xlsx"]          # no temp file left behind


def test_garbage_written_without_an_error_is_never_swapped_in(tmp_path):
    from utility.excel_io import safe_save_workbook

    target = tmp_path / "orders.xlsx"
    _workbook(target, {"Orders": [["Partita Col"], ["A1"]]})
    original = target.read_bytes()

    class Garbage:
        def save(self, destination):
            Path(destination).write_bytes(b"not a workbook")

    with pytest.raises(RuntimeError, match="invalid workbook"):
        safe_save_workbook(Garbage(), target)
    assert target.read_bytes() == original


def test_locked_target_keeps_the_original_and_reports_it(tmp_path, monkeypatch):
    from utility import excel_io

    target = tmp_path / "orders.xlsx"
    _workbook(target, {"Orders": [["Partita Col"], ["A1"]]})
    original = target.read_bytes()
    monkeypatch.setattr(excel_io.os, "replace", lambda *_a: (_ for _ in ()).throw(PermissionError(13, "in use")))
    wb = openpyxl.Workbook()
    with pytest.raises(RuntimeError, match="open or locked"):
        excel_io.safe_save_workbook(wb, target)
    assert target.read_bytes() == original and [p.name for p in tmp_path.iterdir() if p.suffix == ".tmp"] == []


def test_a_save_replaces_the_file_and_leaves_no_previous_copy(tmp_path):
    from utility.excel_io import safe_save_workbook

    target = tmp_path / "orders.xlsx"
    _workbook(target, {"Orders": [["Partita Col"], ["OLD"]]})
    stale = tmp_path / "orders.previous.xlsx"
    stale.write_text("left by an older version", encoding="utf-8")
    wb = openpyxl.Workbook()
    wb.active.append(["NEW"])
    safe_save_workbook(wb, target)
    assert openpyxl.load_workbook(target).active["A1"].value == "NEW"
    assert not stale.exists()                       # and an old leftover is cleaned up


def test_a_previous_copy_is_kept_only_when_asked_for(tmp_path):
    from utility.excel_io import safe_save_workbook

    target = tmp_path / "orders.xlsx"
    _workbook(target, {"Orders": [["Partita Col"], ["OLD"]]})
    wb = openpyxl.Workbook()
    wb.active.append(["NEW"])
    safe_save_workbook(wb, target, keep_backup=True)
    assert openpyxl.load_workbook(tmp_path / "orders.previous.xlsx")["Orders"]["A2"].value == "OLD"


def test_atomic_write_text_keeps_the_old_file_when_the_write_fails(tmp_path, monkeypatch):
    from utility import atomic

    target = tmp_path / "prefs.json"
    target.write_text("old", encoding="utf-8")
    monkeypatch.setattr(atomic.os, "replace", lambda *_a: (_ for _ in ()).throw(OSError("boom")))
    with pytest.raises(OSError):
        atomic.atomic_write_text(target, "new")
    assert target.read_text(encoding="utf-8") == "old" and [p.name for p in tmp_path.iterdir()] == ["prefs.json"]


# ------------------------------------------------------------ 3. parse_number
@pytest.mark.parametrize("value, expected", [
    (1.234, 1.234), (0.125, 0.125), (12.345, 12.345), (3, 3.0), ("12,5", 12.5), ("1,234.56", 1234.56),
    ("1.234,56", 1234.56), ("abc", None), (None, None), (float("nan"), None), (True, None),
])
def test_parse_number_treats_real_numbers_as_numbers(value, expected):
    from utility.utils import parse_number
    assert parse_number(value) == expected


# ------------------------------------------------------------ 2. archive before delete
@pytest.fixture
def history_db(tmp_path, monkeypatch):
    """orders_db functions bound to a throw-away database."""
    from utility import orders_db

    db = tmp_path / "history.sqlite3"
    real = {name: getattr(orders_db, name) for name in ("sync_rows", "archive", "reconcile_sheet")}
    monkeypatch.setattr(orders_db, "sync_rows", lambda p, s, rows: real["sync_rows"](p, s, rows, db_path=db))
    monkeypatch.setattr(orders_db, "archive", lambda p, s, cols, reason="deleted": real["archive"](p, s, cols, reason, db_path=db))
    monkeypatch.setattr(orders_db, "reconcile_sheet",
                        lambda p, s, keys, moved=None: real["reconcile_sheet"](p, s, keys, moved, db_path=db))
    return db


def test_delete_is_refused_when_the_history_cannot_be_written(tmp_path, monkeypatch):
    from exporters import biglietti_exporter
    from utility import orders_db

    path = tmp_path / "shared.xlsx"
    _workbook(path, {"PG-X": [["Partita Col", "Articolo"], ["A1", "C1"]], "Orders": [["Partita Col", "Articolo"]]})
    monkeypatch.setattr(orders_db, "archive", lambda *a, **k: (_ for _ in ()).throw(sqlite3.OperationalError("disk I/O error")))
    monkeypatch.setattr(orders_db, "sync_rows", lambda *a, **k: 0)
    with pytest.raises(RuntimeError, match="nothing was deleted"):
        biglietti_exporter.delete_pg_x_row(path, "A1")
    with pytest.raises(RuntimeError, match="nothing was deleted"):
        biglietti_exporter.delete_shipped_shared_rows(path, ["A1"], "PG-X")
    assert openpyxl.load_workbook(path)["PG-X"]["A2"].value == "A1"          # the row is still there


def test_delete_records_the_history_then_removes_the_row(tmp_path, history_db):
    from exporters import biglietti_exporter
    from utility import orders_db

    path = tmp_path / "shared.xlsx"
    _workbook(path, {"Orders": [["Partita Col", "Articolo"], ["B2", "C2"]], "PG-X": [["Partita Col", "Articolo"]]})
    assert biglietti_exporter.delete_shipped_shared_rows(path, ["B2"], "Orders") == 1
    shipped = orders_db.list_shipped(history_db)
    assert [r["partita_col"] for r in shipped] == ["b2"] and shipped[0]["archived_reason"] == "shipped/invoiced"
    assert openpyxl.load_workbook(path)["Orders"].max_row == 1


def test_manual_delete_removes_multiple_selected_rows_from_orders(tmp_path, history_db):
    from exporters import biglietti_exporter
    from utility import orders_db

    path = tmp_path / "shared.xlsx"
    _workbook(path, {"Orders": [["Partita Col", "Articolo"], ["B2", "C2"], ["B3", "C3"]], "PG-X": [["Partita Col", "Articolo"]]})
    assert biglietti_exporter.delete_shared_rows(path, ["B2", "B3"], "Orders") == 2
    deleted = orders_db.list_history(history_db)
    assert {row["partita_col"] for row in deleted} == {"b2", "b3"}
    assert {row["archived_reason"] for row in deleted} == {"deleted"}
    assert openpyxl.load_workbook(path)["Orders"].max_row == 1


# ------------------------------------------------------------ 5. PG-X -> Orders history
def test_moving_a_row_to_orders_closes_its_pgx_history_record(tmp_path, history_db):
    from exporters import biglietti_exporter
    from utility import orders_db

    path = tmp_path / "shared.xlsx"
    _workbook(path, {"PG-X": [["Partita Col", "Articolo"], ["A1", "C1"]], "Orders": [["Partita Col", "Articolo"]]})
    wb = openpyxl.load_workbook(path)
    biglietti_exporter._sync_workbook_history(path, wb)
    assert {r["sheet_name"] for r in orders_db.list_open(history_db)} == {"PG-X"}

    assert biglietti_exporter.move_pg_x_to_orders(path, "A1") == 1
    open_rows = orders_db.list_open(history_db)
    assert [(r["sheet_name"], r["partita_col"]) for r in open_rows] == [("Orders", "a1")]
    archived = [r for r in orders_db.list_history(history_db) if r["status"] == "archived"]
    assert [(r["sheet_name"], r["archived_reason"]) for r in archived] == [("PG-X", "moved to Orders")]


def test_history_is_not_pruned_when_the_sheet_has_no_partita_col_column(tmp_path, history_db):
    from exporters import biglietti_exporter
    from utility import orders_db

    path = tmp_path / "shared.xlsx"
    _workbook(path, {"Orders": [["Partita Col"], ["A1"]]})
    biglietti_exporter._sync_workbook_history(path, openpyxl.load_workbook(path))
    _workbook(path, {"Orders": [["Something else"], ["A1"]]})
    biglietti_exporter._sync_workbook_history(path, openpyxl.load_workbook(path))
    assert len(orders_db.list_open(history_db)) == 1


# ------------------------------------------------------------ 4 + 6. machines
def test_machine_validation_rules():
    from utility.master_data import validate_machine

    existing = [{"code": "3300", "number": "4", "rocche": "672"}, {"code": "3301", "number": "11", "rocche": "6"}]
    assert validate_machine({"code": "3311", "number": "13", "rocche": "100"}, existing) is None
    assert "Number" in validate_machine({"code": "3311", "number": "abc", "rocche": "100"}, existing)
    assert "Rocche" in validate_machine({"code": "3311", "number": "13", "rocche": "0"}, existing)
    assert "Rocche" in validate_machine({"code": "3311", "number": "13", "rocche": "-5"}, existing)
    assert "already used" in validate_machine({"code": "3311", "number": "4", "rocche": "100"}, existing)
    assert "already used" in validate_machine({"code": "3300", "number": "13", "rocche": "100"}, existing)
    assert validate_machine({"code": "3300", "number": "4", "rocche": "700"}, existing, editing_index=0) is None  # editing itself


def test_bad_machine_rows_never_stop_the_constants_from_loading(tmp_path, monkeypatch):
    from calculate import constants
    from utility import master_data

    monkeypatch.setattr(master_data, "PATH", tmp_path / "master_data.json")
    master_data.PATH.write_text('{"machines": [{"code": "3300", "number": "abc", "rocche": "672"},'
                                ' {"code": "3301", "number": "11", "rocche": "6"}]}', encoding="utf-8")
    try:
        constants.reload_machines()                        # used to raise ValueError
        assert constants.MACHINE_CAPACITY_BY_NUMBER[11] == 6 and 4 in constants.MACHINE_CAPACITY_BY_NUMBER
    finally:
        monkeypatch.undo()
        constants.reload_machines()


def test_saving_a_new_machine_capacity_reaches_the_calculations_without_a_restart(tmp_path, monkeypatch):
    from calculate import abbina_calculator, abbina_suggestions, constants
    from exporters import biglietti_exporter
    from pipelines import ordine_med
    from utility import master_data

    monkeypatch.setattr(master_data, "PATH", tmp_path / "master_data.json")
    data = master_data.load()
    try:
        for machine in data["machines"]:
            if machine["number"] == "7":
                assert machine["rocche"] == "72"
                machine["rocche"] = "80"
        master_data.save(data)
        assert constants.MACHINE_CAPACITY_BY_NUMBER[7] == 80
        assert 80 in abbina_calculator.MACHINE_CAPACITIES and 72 not in abbina_calculator.MACHINE_CAPACITIES
        assert abbina_calculator._smallest_fitting_machine(75) == 80
        assert abbina_suggestions.MACHINE_NUMBERS[80] == 7
        assert biglietti_exporter.MACHINE_SIZE_TABLE[80] == 7
        assert ordine_med.GRUPPO_MACCHINA_TABLE[80] == 3307
    finally:
        monkeypatch.undo()
        constants.reload_machines()
    assert 72 in abbina_calculator.MACHINE_CAPACITIES and 80 not in abbina_calculator.MACHINE_CAPACITIES


# ------------------------------------------------------------ 7. stock grouping
def test_stock_with_the_same_partita_in_two_articles_is_not_merged():
    from calculate.magazino import summarize_by_partita

    frame = pd.DataFrame([
        {"articolo": "G1", "titolo": "30/1", "partita": "P1", "colli": 10, "esistenza": 5.0},
        {"articolo": "G2", "titolo": "40/1", "partita": "P1", "colli": 20, "esistenza": 9.0},
        {"articolo": "G1", "titolo": "30/1", "partita": "P1", "colli": 2, "esistenza": 1.0},
    ])
    out = summarize_by_partita(frame).set_index("articolo")
    assert len(out) == 2
    assert (out.loc["G1", "mag_rocche"], out.loc["G2", "mag_rocche"]) == (12, 20)


# ------------------------------------------------------------ 8. disk cache
def test_cache_file_names_are_valid_on_windows_and_do_not_collide(tmp_path, monkeypatch):
    from utility import disk_cache

    monkeypatch.setattr(disk_cache, "APP_DATA_DIR", tmp_path)
    key = ("file.xlsx", 1, 2)
    first = disk_cache._cache_path("magazino:None", key)
    second = disk_cache._cache_path("magazino:('G130', 'G170')", key)
    assert first != second
    for path in (first, second):
        assert not set(':*?"<>|') & set(path.name)


def test_disk_cache_really_writes_a_file_and_hands_out_private_copies(tmp_path, monkeypatch):
    from utility import disk_cache

    monkeypatch.setattr(disk_cache, "APP_DATA_DIR", tmp_path / "app")
    source = tmp_path / "stock.xlsx"
    source.write_text("x", encoding="utf-8")
    disk_cache.invalidate()
    calls = []

    def compute():
        calls.append(1)
        return pd.DataFrame({"a": [1, 2]}), []

    first, _ = disk_cache.cached_load("magazino:None", source, compute)
    assert list((tmp_path / "app" / "cache").glob("*.pkl"))                      # the disk copy exists now
    first.loc[0, "a"] = 999                                                       # the caller edits its copy ...
    second, _ = disk_cache.cached_load("magazino:None", source, compute)
    assert second.loc[0, "a"] == 1 and len(calls) == 1                           # ... the cache is unaffected
    disk_cache.invalidate()
    third, _ = disk_cache.cached_load("magazino:None", source, compute)          # served from the disk file
    assert third.loc[0, "a"] == 1 and len(calls) == 1


# ------------------------------------------------------------ backups
def test_backup_skips_dev_folders_and_snapshots_sqlite(tmp_path, monkeypatch):
    from utility import backup

    app = tmp_path / "app"
    (app / "data").mkdir(parents=True)
    (app / "venv" / "lib").mkdir(parents=True)
    (app / "venv" / "lib" / "big.py").write_text("x", encoding="utf-8")
    (app / "settings").mkdir()
    (app / "settings" / "prefs.json").write_text("{}", encoding="utf-8")
    db = app / "data" / "orders.sqlite3"
    conn = sqlite3.connect(db)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("CREATE TABLE t(x)")
    conn.execute("INSERT INTO t VALUES (42)")
    conn.commit()                                  # the row is in the -wal file, not yet in the main file
    monkeypatch.setattr(backup, "APP_DATA_DIR", app)
    try:
        archive_path = backup.create_backup(tmp_path / "planning_backup_test.zip")
    finally:
        conn.close()
    with zipfile.ZipFile(archive_path) as archive:
        names = set(archive.namelist())
        assert "settings/prefs.json" in names and "data/orders.sqlite3" in names
        assert not any(n.startswith("venv/") or n.endswith(("-wal", "-shm")) for n in names)
        archive.extract("data/orders.sqlite3", tmp_path / "out")
    assert sqlite3.connect(tmp_path / "out" / "data" / "orders.sqlite3").execute("SELECT x FROM t").fetchall() == [(42,)]


def test_restore_removes_stale_wal_files_of_the_replaced_database(tmp_path, monkeypatch):
    from utility import backup

    app = tmp_path / "app"
    (app / "data").mkdir(parents=True)
    (app / "data" / "orders.sqlite3-wal").write_bytes(b"stale wal")
    (app / "data" / "orders.sqlite3-shm").write_bytes(b"stale shm")
    zip_path = tmp_path / "b.zip"
    good = tmp_path / "good.sqlite3"
    sqlite3.connect(good).execute("CREATE TABLE t(x)").connection.commit()
    with zipfile.ZipFile(zip_path, "w") as archive:
        archive.write(good, "data/orders.sqlite3")
    monkeypatch.setattr(backup, "APP_DATA_DIR", app)
    assert backup.restore_backup(zip_path) == 1
    assert not (app / "data" / "orders.sqlite3-wal").exists() and not (app / "data" / "orders.sqlite3-shm").exists()


# ------------------------------------------------------------ build environment
def test_venv_prune_keeps_dev_tools_and_their_dependencies():
    import sync_venv_packages as sync

    kept = sync._with_dependencies({"pytest"})
    assert {"pytest", "pluggy", "iniconfig"} <= kept


# ------------------------------------------------------------ one writer at a time
def test_shared_workbook_writers_never_run_at_the_same_time():
    import threading
    import time

    from exporters.biglietti_exporter import _serialized_workbook_write

    active, overlaps = [], []

    @_serialized_workbook_write
    def writer():
        active.append(1)
        if len(active) > 1:
            overlaps.append(1)
        time.sleep(0.02)
        active.pop()

    threads = [threading.Thread(target=writer) for _ in range(6)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    assert overlaps == []


# ------------------------------------------------------------ EXE must contain lazily loaded modules
def test_pyinstaller_spec_collects_every_package_that_is_loaded_lazily_by_name():
    """parsers.kamal_parser etc. are only ever imported by name at runtime, so
    PyInstaller cannot see them: they were missing from the EXE and Ordine Kamal
    failed with "No module named 'parsers.kamal_parser'".  The spec must collect
    every first-party package, and every lazily referenced module must live in one."""
    import re

    root = Path(__file__).resolve().parent.parent
    spec = (root / "main.spec").read_text(encoding="utf-8")
    match = re.search(r"FIRST_PARTY_PACKAGES\s*=\s*\(([^)]*)\)", spec)
    assert match, "main.spec must define FIRST_PARTY_PACKAGES"
    collected = set(re.findall(r"[\"']([A-Za-z_][\w]*)[\"']", match.group(1)))
    assert "collect_submodules(_pkg)" in spec

    packages = {p.name for p in root.iterdir() if (p / "__init__.py").is_file() and p.name != "tests"}
    assert packages <= collected, f"packages not collected into the EXE: {sorted(packages - collected)}"

    referenced = set()
    for path in root.rglob("*.py"):
        if "tests" in path.parts or ".git" in path.parts:
            continue
        text = path.read_text(encoding="utf-8", errors="ignore")
        referenced |= set(re.findall(r"(?:lazy_call|LazyModule|import_module)\(\s*[\"']([\w.]+)[\"']", text))
    for module in referenced:
        top = module.split(".")[0]
        assert top in collected, f"{module} is loaded by name but {top!r} is not collected in main.spec"
        assert (root / Path(*module.split("."))).with_suffix(".py").is_file() or (root / Path(*module.split(".")) / "__init__.py").is_file(), module


# ------------------------------------------------------------ EL KAMAL (Cliente 3019) in the ERP dump
ERP_HEADERS = ["Cliente", " ", "Articolo", " 2", " 3", "Colore", " 4", "Sottolinea", "Confezione", "Ordine", "Riga",
               "Tipo riga Ordine", "Prezzo", "Data", "Consegna", " 5", "Riserva", "Unità di misura", "Partita",
               "Ordinata", "Spedita", "Rimanenza", "Chili Disposti Scatola", "Assegnata", "Scatole assegnate",
               "Scatole spedibili", " 6", "Disposta", "Saldato", "Descrizione aggiuntiva ordine"]


def _kamal_erp_dump(path: Path, rows: list[tuple]) -> None:
    sheet_rows = [ERP_HEADERS, ERP_HEADERS]            # the real export repeats its header row
    for code, article, colore, riga, partita, qty, comment in rows:
        row = [None] * len(ERP_HEADERS)
        row[0], row[1], row[2], row[5] = code, "AL KAMAL INDUSTRY AND TRADE", article, colore
        row[6], row[9], row[10] = f"KA-{colore}", "0/005420", riga
        row[14], row[17], row[18] = datetime(2026, 11, 3), "KG", partita
        row[19], row[29] = f"   {qty},00", comment
        sheet_rows.append(row)
    _workbook(path, {"Sheet1": sheet_rows})


def _dispo_file(path: Path, rows: list[tuple]) -> None:
    header = ["Dispo", "Field1", "Field2", "Field3", "Articolo", "Colore", "Field5", "Field8", "Field9", "Field10", "Field11", "Field14"]
    _workbook(path, {"Sheet1": [header, header, *rows]})


def test_erp_dump_with_cliente_3019_is_read_as_el_kamal(tmp_path):
    from exporters import biglietti_exporter as b

    data, dispo = tmp_path / "ORD.xlsx", tmp_path / "Dispo.xlsx"
    _kamal_erp_dump(data, [(3019, "C170026S", 790001, 1, 159527, 672, "PG-159392-ORD.nr.25-2026(LETTER 105)"),
                           (3019, "C170027S", 790129, 2, 159528, 84, "PG-159393-ORD.nr.25-2026(LETTER 105)")])
    _dispo_file(dispo, [("D-00505578-001", 79185, "T1303", 3309, "C170026S", "790001", "KA-OPTICAL", "OO", 698, 88, "KG", "003019"),
                        ("D-00505578-002", 79186, "T1304", 3308, "C170027S", "790129", "KA-BLE 9", "OO", 87, 36, "KG", "003019")])
    assert b.detect_order_format(data) == "ELVY_MED"      # same dump shape as ELVY / MED
    records, _raw = b.load_order(data, dispo)
    assert [r.customer_code for r in records] == ["3019", "3019"]
    first, second = records
    assert (first.bagno, first.machine, first.quantity_cones, first.kg) == ("T1303", "3", 88, 69.8)     # ERP tenths: 698 -> 69.8
    assert (second.bagno, second.machine, second.quantity_cones, second.kg) == ("T1304", "5", 36, 8.7)  # ERP tenths: 87 -> 8.7
    assert first.raw_batch == "159392" and first.colored_batch == "159527"
    assert first.commento == "PG-159392-ORD.nr.25-2026(LETTER 105)"
    assert b.build_el_kamal_stem(records) == "D-505578-KAMAL-2026"
    b.enrich_records(records, "EL_KAMAL")
    assert (first.machine, first.kg) == ("3", 69.8)                         # enrichment leaves Kamal's own values alone
    assert b._ticket_header(first, "x").startswith("EL KAMAL - ")


def test_customer_name_for_3019_comes_from_master_data_customers(tmp_path, monkeypatch):
    from exporters import biglietti_exporter as b
    from utility import master_data

    monkeypatch.setattr(master_data, "PATH", tmp_path / "master_data.json")
    data = master_data.load()
    data["customers"] = [{"code": "3019", "name": "EL KAMAL TEXTILE"}]
    master_data.save(data)
    assert b._master_data_customer_name("3019", "EL KAMAL") == "EL KAMAL TEXTILE"
    assert b._master_data_customer_name("9999", "fallback") == "fallback"


def test_kamal_row_without_a_dispo_match_keeps_the_order_kg_and_no_guessed_machine(tmp_path):
    from exporters import biglietti_exporter as b

    data = tmp_path / "ORD.xlsx"
    _kamal_erp_dump(data, [(3019, "C170026S", 790001, 1, 159527, 672, "PG-1-ORD.nr.25-2026(LETTER 105)")])
    (record,), _ = b.load_order(data, None)
    assert record.kg == 672 and record.machine == "" and record.quantity_cones is None and record.bagno == ""
    assert b._machine_number_from_dispo_code("9999") == "" and b._machine_number_from_dispo_code("abc") == ""


def test_elvy_and_med_rows_are_still_read_next_to_kamal_rows(tmp_path):
    from exporters import biglietti_exporter as b

    data = tmp_path / "ORD.xlsx"
    _kamal_erp_dump(data, [(3009, "C130027S", 440, 1, 1001, 24, "PG-1-PO-5-2026"),
                           (3004, "C130028S", 441, 2, 1002, 24, "PG-2-PO-6-2026"),
                           (3019, "C170026S", 790001, 3, 1003, 672, "PG-3-ORD.nr.25-2026(LETTER 105)"),
                           (3011, "C130029S", 442, 4, 1004, 24, "ignored customer")])
    records, _ = b.load_order(data, None)
    assert [r.customer_code for r in records] == ["3009", "3004", "3019"]


def test_biglietti_tab_routes_3019_rows_to_the_el_kamal_outputs():
    root = Path(__file__).resolve().parent.parent
    tab = (root / "gui" / "tabs" / "biglietti_tab.py").read_text(encoding="utf-8")
    assert 'has_kamal = any(r.customer_code == "3019"' in tab
    assert "kamal_records = [r for r in records if r.customer_code == \"3019\"]" in tab
    assert 'customer="EL_KAMAL"' in tab


def test_kamal_colour_name_is_the_cell_after_the_colour_code_in_the_order_file(tmp_path):
    """The (longer) name next to the colour code in the uploaded order file is
    used -- the Dispo file's shorter label must never replace it."""
    from exporters import biglietti_exporter as b

    data, dispo = tmp_path / "ORD.xlsx", tmp_path / "Dispo.xlsx"
    _kamal_erp_dump(data, [(3019, "C170115S", 790351, 1, 159527, 330, "PG-1-ORD.Nr.28-2026(LETTER 108)")])
    ws_rows = [list(r) for r in openpyxl.load_workbook(data).active.iter_rows(values_only=True)]
    for r in ws_rows[2:]:
        r[6] = "KA-MENT 4 - 13-0608 TCX       "
    _workbook(data, {"Sheet1": ws_rows})
    _dispo_file(dispo, [("D-00505582-001", 79228, "T1343", 3306, "C170115S", "790351", "KA-MENT 4", "OO", 33, 28, "KG", "003019")])
    (rec,), _ = b.load_order(data, dispo)
    assert rec.color_name == "KA-MENT 4 - 13-0608 TCX"


def test_kamal_colour_name_follows_the_colour_code_column_when_columns_shift(tmp_path):
    from exporters import biglietti_exporter as b

    data = tmp_path / "ORD.xlsx"
    _kamal_erp_dump(data, [(3019, "C170115S", 790351, 1, 159527, 330, "PG-1-ORD.Nr.28-2026(LETTER 108)")])
    rows = [list(r) for r in openpyxl.load_workbook(data).active.iter_rows(values_only=True)]
    rows[2][6] = "KA-MENT 4 - 13-0608 TCX"
    shifted = [[None] + r for r in rows]                     # one extra leading column
    _workbook(data, {"Sheet1": shifted})
    (rec,), _ = b.load_order(data, None)
    assert rec.color_code == "790351" and rec.color_name == "KA-MENT 4 - 13-0608 TCX"
