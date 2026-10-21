"""
kamal_excel_exporter.py
Builds the Kamal order Excel workbook: a raw "Kamal Order" sheet (parsed
PDF rows, including Dye Type), an "Ordine Kamal" sheet (matched rows,
mirroring Ordini ELVY's structure), and -- if raw yarn stock was supplied
-- a "Filato x Tinturia" sheet listing what was assigned.
"""
from __future__ import annotations
from utility.excel_io import safe_save_workbook

from pathlib import Path
from typing import Sequence

import openpyxl
from utility.utils import clean_text
from utility.master_data import raw_articolo_for
from parsers.dfm_lookup import is_first_time_dyeing

from exporters.excel_exporter import (
    FILATO_TINTURIA_COLUMNS,
    apply_column_widths,
    freeze_and_filter,
    write_data,
    write_header,
)
from parsers.kamal_parser import KamalOrderRow
from pipelines.ordine_kamal import (
    OrdineKamalRow,
    assign_bagno_prepostato,
    assign_ordine_kamal_machines,
    build_ordine_kamal_rows,
    match_by_lotto,
    resolve_kamal_colours,
)
from pipelines.ordini_elvy import RawYarnMatch, match_raw_yarn
from utility.utils import logger

RAW_KAMAL_COLUMNS: list[tuple[str, str, str]] = [
    ("Reply No",       "reply_no",       "text"),
    ("Order Date",     "order_date",     "text"),
    ("Dye Type",       "dye_type",       "text"),
    ("Notes",          "notes",          "text"),
    ("Sender Name",    "sender_name",    "text"),
    ("Sender Msg No",  "sender_msg_no",  "text"),
    ("COLOREDFM",      "coloredfm",      "text"),
    ("Titolo",         "titolo",         "text"),
    ("Colore",         "colore_name",    "text"),
    ("Peso (kg)",      "peso_kg",        "number"),
    ("Abbina",         "abbina",         "text"),
    ("Check Articolo", "check_articolo", "text"),
]

ORDINE_KAMAL_COLUMNS: list[tuple[str, str, str]] = [
    ("CLIENTE",              "cliente",              "integer"),
    ("ARTICOLO",             "articolo_delta",       "text"),
    ("COLORE",               "coloredfm",            "integer"),
    ("Q.TA",                 "peso_kg",              "number"),
    ("CONSEGNA",             "data_consegna",        "date"),
    ("COMMENTO",             "commento",             "text"),
    ("LAVORANTE",            "lavorante",            "integer"),
    ("LAV. SUCC",            "lavorante_successivo", "integer"),
    ("DATA RICONSEGNA",      "data_riconsegna",      "date"),
    ("MAG. GREGGIO",         "stock_mag",            "integer"),
    ("SIGLA DISPOSIZIONE",   "sigla_dispo",          "text"),
    ("BAGNO PREPOSTO",       "bagno_proposto",       "text"),
    ("GRUPPO MACCHINA",      "macchina",             "integer"),
]


class KamalExcelExporter:
    """Writes the Kamal workbook from the parsed Kamal order PDF rows.

    Sheets: "Kamal Order" (the PDF rows as read, with the DFM colour code/name and
    the Abbina / first-time-dyeing checks filled in), "Ordine Kamal" (the rows
    prepared for the ERP: Articolo, machine group, bagno, raw yarn) and
    "Filato x Tinturia" (the raw yarn to prepare, written when there is some).  Inputs besides the rows: the DFM C170 entries, the magazino and lotti
    summaries, the Articoli codes map, the first bagno and the order number.
    """
    def __init__(self, rows: Sequence[KamalOrderRow], output_path: Path,
                 dfm_c170_entries: list[dict] | None = None,
                 magazino_summary=None, codes_map: dict | None = None,
                 lotti_summary=None, first_bagno: str | None = None,
                 order_no: str | None = None) -> None:
        self.rows = rows
        self.output_path = output_path
        self.dfm_c170_entries = dfm_c170_entries or []
        self.magazino_summary = magazino_summary
        self.codes_map = codes_map
        self.lotti_summary = lotti_summary
        # "First bagno" for BAGNO PREPOSTO (None / "" = leave blank).
        self.first_bagno = first_bagno
        self.order_no = order_no

    def export(self) -> None:
        """Build and save the workbook.

        Order of work: look every colour up in DFM (so every sheet shows the DFM code
        and name), build the Ordine Kamal rows, match raw yarn (first by the supplier's
        Lotto message number, then by quantity against the magazino), assign the machine
        group and bagno, copy Abbina and the first-time-dyeing check back to the PDF
        rows, then write the sheets.
        """
        logger.info("Exporting %d Kamal row(s) → %s", len(self.rows), self.output_path.name)

        # COLOREDFM and the colour name come from DFM where it knows the colour;
        # done before the raw sheet is written so that sheet shows them too.
        resolve_kamal_colours(self.rows, self.dfm_c170_entries)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Kamal Order"
        ws.sheet_view.rightToLeft = False
        write_header(ws, RAW_KAMAL_COLUMNS)
        write_data(ws, RAW_KAMAL_COLUMNS, self.rows)
        apply_column_widths(ws, RAW_KAMAL_COLUMNS)
        freeze_and_filter(ws, RAW_KAMAL_COLUMNS)

        kamal_rows = build_ordine_kamal_rows(
            list(self.rows), self.dfm_c170_entries, order_no=self.order_no
        )

        # Primary: match by Kamal's own raw-yarn message number (Lotto),
        # which is a deterministic reference from their supplier.
        if self.lotti_summary is not None and not self.lotti_summary.empty:
            lotto_matched = match_by_lotto(kamal_rows, self.lotti_summary)
            logger.info("Raw yarn matched via Lotto: %d row(s)", lotto_matched)

        raw_yarn_matches = self._build_lotto_matches(kamal_rows)

        if self.magazino_summary is not None and not self.magazino_summary.empty:
            fallback_matches = match_raw_yarn(
                kamal_rows, self.magazino_summary, self.codes_map, quantity_attr="peso_kg"
            )
            if fallback_matches:
                raw_yarn_matches.extend(fallback_matches)
            logger.info("Raw yarn matched via quantity fallback: %d batch(es) assigned", len(fallback_matches))

        ws2 = wb.create_sheet("Ordine Kamal")
        ws2.sheet_view.rightToLeft = False
        # Use the same machine calculation for the workbook and ERP update.
        assign_ordine_kamal_machines(kamal_rows)
        assign_bagno_prepostato(kamal_rows, self.first_bagno)

        # Also set Abbina on the original parsed rows (first sheet).  kamal_rows
        # is built 1:1, in order, from self.rows, so copy position by position:
        # this also covers rows grouped by Colore (blank COLOREDFM), which a
        # COLOREDFM lookup used to leave without an Abbina.
        for parsed_row, derived_row in zip(self.rows, kamal_rows):
            parsed_row.abbina = derived_row.abbina or ""
            # Use ELVY's exact Articolo + COLOREDFM history rule.  Kamal's
            # derived Articolo is finished (C170...), so resolve it to its
            # raw equivalent before passing it to the shared helper.
            parsed_row.check_articolo = (
                "prima volta tint."
                if is_first_time_dyeing(
                    raw_articolo_for(derived_row.articolo_delta),
                    derived_row.coloredfm,
                    self.dfm_c170_entries,
                )
                else ""
            )

        # The raw sheet was written before the derived rows were calculated.
        # Update its Abbina cells after the machine assignment so the parsed
        # Kamal Order sheet contains the same values used by Ordine Kamal.
        abbina_col = next(
            (index + 1 for index, (_header, attr, _dtype) in enumerate(RAW_KAMAL_COLUMNS)
             if attr == "abbina"),
            None,
        )
        if abbina_col is not None:
            for row_index, parsed_row in enumerate(self.rows, start=2):
                ws.cell(row=row_index, column=abbina_col).value = getattr(parsed_row, "abbina", "") or ""

        check_articolo_col = next(
            (index + 1 for index, (_header, attr, _dtype) in enumerate(RAW_KAMAL_COLUMNS)
             if attr == "check_articolo"),
            None,
        )
        if check_articolo_col is not None:
            for row_index, parsed_row in enumerate(self.rows, start=2):
                ws.cell(row=row_index, column=check_articolo_col).value = (
                    getattr(parsed_row, "check_articolo", "") or ""
                )

        write_header(ws2, ORDINE_KAMAL_COLUMNS)
        write_data(ws2, ORDINE_KAMAL_COLUMNS, kamal_rows)
        apply_column_widths(ws2, ORDINE_KAMAL_COLUMNS)
        freeze_and_filter(ws2, ORDINE_KAMAL_COLUMNS)

        if raw_yarn_matches:
            ws3 = wb.create_sheet("Filato x Tinturia")
            ws3.sheet_view.rightToLeft = False
            write_header(ws3, FILATO_TINTURIA_COLUMNS)
            write_data(ws3, FILATO_TINTURIA_COLUMNS, raw_yarn_matches)
            apply_column_widths(ws3, FILATO_TINTURIA_COLUMNS)
            freeze_and_filter(ws3, FILATO_TINTURIA_COLUMNS)

            from openpyxl.styles import Alignment, Font
            label_col_idx = next(
                (i + 1 for i, (_, attr, _) in enumerate(FILATO_TINTURIA_COLUMNS) if attr == "label"), None
            )
            if label_col_idx:
                bold_label_font = Font(bold=True, name="Calibri", size=11)
                for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row,
                                          min_col=label_col_idx, max_col=label_col_idx):
                    for cell in row:
                        cell.font = bold_label_font
                        cell.alignment = Alignment(horizontal="center", vertical="center")

        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        safe_save_workbook(wb, self.output_path, keep_backup=False)
        logger.info("Export completed: %s", self.output_path)

    def _build_lotto_matches(self, kamal_rows: list[OrdineKamalRow]) -> list[RawYarnMatch]:
        """Build the raw-yarn matches from the Partita written in each row's Commento.
        Rows marked PG-X or with an empty/X partita are skipped; the quantities already assigned per (article, titolo, partita) are accumulated.
        """
        matches: list[RawYarnMatch] = []
        if self.lotti_summary is None or self.lotti_summary.empty:
            return matches

        assigned: dict[tuple[str, str, str], float] = {}
        for row in kamal_rows:
            commento = clean_text(row.commento)
            if "PG-X" in commento.upper():
                continue
            partita = self._extract_partita(commento)
            if not partita or partita.upper() == "X":
                continue
            articolo = raw_articolo_for(row.articolo_delta) if row.articolo_delta.upper().startswith("C") else row.articolo_delta
            titolo = (self.codes_map or {}).get(row.articolo_delta, "")
            key = (articolo, titolo, partita)
            assigned[key] = assigned.get(key, 0.0) + (row.peso_kg or 0.0)

        for (articolo, titolo, partita), peso in assigned.items():
            matches.append(RawYarnMatch(
                articolo=articolo,
                titolo=titolo,
                partita=partita,
                rocce=round(peso, 2),
                peso=round(peso, 2),
            ))
        return matches

    @staticmethod
    def _extract_partita(commento: str) -> str:
        if not commento:
            return ""
        normalized = commento.upper()
        idx = normalized.find("PG-")
        if idx == -1:
            return ""
        start = idx + 3
        end = start
        while end < len(commento) and commento[end] not in (" ", "\t", "\r", "\n", "-", "("):
            end += 1
        return commento[start:end].strip()
