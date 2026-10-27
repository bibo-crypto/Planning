"""The "Prezzi" price-check workbook, shared by Ordine MED and Ordine Elvy.

Single responsibility: lay out and style the sheet.  What goes in the rows is decided
by each order page (MED rows come from the MED order file, Elvy rows from the PO PDF);
both hand over plain lists in PREZZI_HEADERS order.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable, Sequence

from utility.excel_io import safe_save_workbook

# The column order the planner works with (same as the reference sheet).
PREZZI_HEADERS = [
    "Riga", "ARTICOLO", "TITOLO", "COLORE", "Q.TA", "DESCR COL",
    "LIVELLO", "M/C", "PREZZO", "prezzo + (2$)",
]


def price_cell(value: Any) -> Any:
    """A price that could not be found stays an empty cell instead of a blank string."""
    return None if value in (None, "") else value


def write_prezzi_workbook(path: Path, rows: Iterable[Sequence[Any]], sheet_title: str = "MED") -> None:
    """Write ``path``: one sheet (``sheet_title``), with PREZZI_HEADERS and one line per row.

    Styled like the planner's reference sheet: navy headers for the order columns,
    red headers for the two price columns, LIVELLO shaded red, prices with two
    decimals, and a colour code that appears more than once highlighted in green.
    """
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Font, PatternFill
    from exporters.biglietti_exporter import _style_sheet

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(PREZZI_HEADERS)
    for row in rows:
        ws.append(list(row))
    _style_sheet(ws)

    navy, dark, red = "1F2A5C", "000000", "C00000"
    pink = "F4B6B6"
    header_fills = {
        "Riga": navy, "ARTICOLO": navy, "TITOLO": navy, "COLORE": navy, "Q.TA": navy,
        "DESCR COL": dark, "LIVELLO": pink, "M/C": dark, "PREZZO": red, "prezzo + (2$)": red,
    }
    for cell in ws[1]:
        fill = header_fills[cell.value]
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(bold=True, color=red if fill == pink else "FFFFFF")
    for column, name in enumerate(PREZZI_HEADERS, start=1):
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=column)
            if name == "LIVELLO":
                cell.fill = PatternFill("solid", fgColor=pink)
                cell.font = Font(color=red)
            elif name in ("PREZZO", "prezzo + (2$)"):
                cell.number_format = "0.00"
        longest = max((len(str(ws.cell(row=row, column=column).value or "")) for row in range(2, ws.max_row + 1)), default=0)
        ws.column_dimensions[ws.cell(row=1, column=column).column_letter].width = max(12, len(name) + 4, longest + 3)

    last = ws.max_row
    if last > 1:
        colore = ws.cell(row=1, column=PREZZI_HEADERS.index("COLORE") + 1).column_letter
        green = PatternFill(start_color="FFC6EFCE", end_color="FFC6EFCE", fill_type="solid")
        ws.conditional_formatting.add(
            f"{colore}2:{colore}{last}",
            FormulaRule(
                formula=[f'AND(${colore}2<>"",COUNTIF(${colore}$2:${colore}${last},{colore}2)>1)'],
                fill=green,
            ),
        )
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    safe_save_workbook(wb, path, keep_backup=False)
    wb.close()
